#!/usr/bin/env python3
"""
Generates the sample data for the agents lightning lesson.

WHAT THIS DATA IS FOR
---------------------
The demo agent reads one folder and writes FOUNDER_BRIEF.md. The folder holds
ONE file: a week of customer messages exported from a small business's support
inbox, as a spreadsheet. That is the honest shape of "rung one" -- no email
integration, no helpdesk API. The owner exported last week and dropped the file
in a folder. One source, one row per message, with the metadata a real export
carries (timestamp, channel, customer, plan).

The business is Bookable, a small scheduling + invoicing SaaS for service
businesses (cleaners, trainers, tutors). That makes a QuickBooks sync request
natural, and it matches the message set already on the deck's slides.

THE SEEDED SET (the deck depends on these -- keep them if you regenerate)
-------------------------------------------------------------------------
- Two clear escalations: a DOUBLE CHARGE (money, wrong right now) and a
  CANCELLATION on the top plan whose stated reason is a missing feature (data
  export), not price.
- Two clear drops: SPAM and a THANK-YOU note.
- One ordinary support issue: a LOGIN problem. Handled, not escalated. Do NOT
  add a follow-up from the same customer -- that would push the agent to
  escalate it and contradict the "Read the brief" slide.
- THREE versions of the same question, worded differently, from three
  different customers: does it work with QuickBooks. No single row contains
  the finding; it only exists across the sheet. This is the beat that makes
  the session.

Around those, realistic filler a week of inbox actually has: an out-of-office
auto-reply, a "any update?" follow-up on a feature request, a vendor pitch, a
password reset, a refund policy question, a team-pricing inquiry from a
prospect. These give the agent more to drop, and one thing (the prospect) it
may reasonably flag as revenue rather than drop -- a judgment call worth
having on screen.

WHERE THE FILE GOES
-------------------
    agents/sample-data/inbox/support-export_2026-09-08_to_2026-09-14.xlsx

Point Cowork at agents/sample-data/inbox/ -- NOT at agents/sample-data/. If
Claude can see this script it will read the seeded design and the demo is
spoiled. The inbox folder must contain only the export.

Run: python3 generate.py   (from this directory; needs openpyxl)
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT_DIR = Path(__file__).parent / "inbox"
OUT_FILE = OUT_DIR / "support-export_2026-09-08_to_2026-09-14.xlsx"

COLUMNS = ["Ticket", "Received", "Channel", "Name", "Email", "Plan", "Subject", "Message", "Status"]

# (ticket, received, channel, name, email, plan, subject, message)
ROWS = [
    (4817, "2026-09-08 08:42", "Contact form", "Priya Natarajan", "priya@brightstepcleaning.com", "Starter",
     "Can't log in",
     "I can't log into my account this morning. It says my password is wrong but I haven't changed it. "
     "I have three jobs today and need the schedule."),

    (4818, "2026-09-08 10:15", "Email", "Marcus Whitfield", "marcus@whitfieldpt.com", "Pro",
     "QuickBooks?",
     "Quick question before I commit to the annual plan -- do you integrate with QuickBooks? "
     "I'd like invoices to land in there automatically."),

    (4819, "2026-09-08 19:03", "In-app chat", "Dana Kowalski", "dana.k@gmail.com", "Starter",
     "Feature idea",
     "Would love a dark mode. I do my scheduling at night and the white screen is a lot. Not urgent, just a wish!"),

    (4820, "2026-09-09 07:56", "Email", "Tom Reyes", "tom@reyeselectric.net", "Pro",
     "Charged twice",
     "Your checkout charged me twice for September. I have two $49 charges on my card from Bookable dated the 8th. "
     "Please sort this out, I need one of them back."),

    (4821, "2026-09-09 11:30", "Contact form", "Leah Brennan", "leah@brennanconsulting.co", "Pro",
     "Affiliate program",
     "I recommend Bookable to a lot of my consulting clients. Can you add an affiliate program? "
     "Happy to be a test case."),

    (4822, "2026-09-09 16:47", "In-app chat", "Sam Okafor", "sam.okafor@outlook.com", "Starter",
     "Thanks",
     "You guys rock. Switched from a paper diary in March and I've not missed a booking since. That's all, carry on."),

    (4823, "2026-09-10 09:12", "Email", "Rachel Lindqvist", "rachel@lindqvisttutoring.com", "Business",
     "Cancelling my subscription",
     "I'm cancelling at the end of this cycle. It's not the price -- I can't export my data. I have two years of "
     "client history in here and no way to get it out. If that changes let me know, I'd come back."),

    (4824, "2026-09-10 13:25", "Contact form", "Jordan Ellis", "jordan@ellisandsons.com", "Business",
     "Accounting",
     "Does this work with QuickBooks Online? My accountant is asking whether she can pull invoices from Bookable "
     "directly or whether I have to keep sending her PDFs."),

    (4825, "2026-09-10 14:02", "Email", "Bookable Growth Team", "listings@biz-verify-now.com", "",
     "URGENT: claim your business listing now",
     "Your business listing is UNVERIFIED and may be removed within 48 hours. Click here to claim your listing "
     "and secure your position in local search results. Act now."),

    (4826, "2026-09-11 08:20", "In-app chat", "Ana Ferreira", "ana@ferreirapetcare.com", "Pro",
     "Sync",
     "Any chance of a QuickBooks sync? I'm doing double entry every month end and it takes me a whole evening."),

    (4827, "2026-09-11 08:21", "Email", "Ana Ferreira", "ana@ferreirapetcare.com", "Pro",
     "Automatic reply: Sync",
     "Thanks for your email. I'm out on jobs until 5pm and will reply this evening. For bookings please use the "
     "online form. -- Ana"),

    (4828, "2026-09-11 15:40", "Contact form", "Chris Nakamura", "chris@nakamurafitness.com", "Starter",
     "Following up",
     "Following up on my request from August for recurring invoices -- any timeline on that? Not a dealbreaker, "
     "just planning ahead."),

    (4829, "2026-09-12 10:05", "Email", "Ben at LocalBiz Weekly", "ben@localbizweekly.com", "",
     "Feature Bookable in our newsletter",
     "Hi there! We'd love to feature Bookable in our small business newsletter (14k subscribers). Sponsored "
     "placement starts at $500. Let me know if you'd like our media kit."),

    (4830, "2026-09-12 17:33", "In-app chat", "Maya Thornton", "maya.thornton@yahoo.com", "Starter",
     "Password",
     "Can you reset my password? I've lost access to the email I signed up with. New email is this one."),

    (4831, "2026-09-13 09:48", "Contact form", "Derek Oyelaran", "derek@oyelaranlandscapes.com", "Pro",
     "Refund for August?",
     "I didn't use the app at all in August (was away). Is there any chance of a refund or credit for that month? "
     "Understand if not."),

    (4832, "2026-09-14 11:19", "Contact form", "Nicole Vasquez", "nicole@vasquezhomeservices.com", "",
     "Plan for a team of 12?",
     "We're a home services company with 12 field staff and two office admins. Do you have a plan for teams that "
     "size, and can admins see everyone's calendar? Looking to decide this month."),
]


def build() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tickets"

    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EDEDED")
        cell.alignment = Alignment(vertical="top")

    for ticket, received, channel, name, email, plan, subject, message in ROWS:
        dt = datetime.strptime(received, "%Y-%m-%d %H:%M")
        ws.append([f"#{ticket}", dt, channel, name, email, plan, subject, message, "Open"])

    # the way a helpdesk export actually looks: dates formatted, message column wide and wrapped
    widths = {"A": 8, "B": 17, "C": 13, "D": 20, "E": 30, "F": 9, "G": 30, "H": 80, "I": 8}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "yyyy-mm-dd hh:mm"
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(ROWS) + 1}"

    OUT_DIR.mkdir(exist_ok=True)
    wb.save(OUT_FILE)
    return OUT_FILE


if __name__ == "__main__":
    path = build()
    print(f"wrote {path.relative_to(Path(__file__).parent.parent.parent)}  ({len(ROWS)} rows)")
